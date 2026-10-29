"""Lectura y validación del maestro de accesorios (hoja "Maestro"; hoja opcional "Modelos")."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import pandas as pd

from .config import PRECIO_MARCADOR
from .tipos import AVISO, ERROR, INFO, Hallazgo
from .util import clave, stem

TIPOS = ("Accesorio", "Color Studio", "Mantenimiento")
OBLIGATORIAS = ("codigo", "tipo", "nombre", "precio_usd", "controla_stock", "activo")
OPCIONALES = (
    "codigo_sistema", "descripcion", "categoria", "modelos", "imagen", "orden", "precio_origen", "observaciones",
)


class MaestroInvalido(Exception):
    """El maestro no tiene el formato esperado."""


@dataclass
class Maestro:
    items: pd.DataFrame
    modelos: dict[str, str]   # código de modelo -> nombre comercial (puede estar vacío)


def _si_no(valor) -> bool | None:
    s = str(valor).strip().upper() if valor is not None else ""
    if s in ("SI", "SÍ", "S", "1", "TRUE", "X", "VERDADERO"):
        return True
    if s in ("NO", "N", "0", "FALSE", "FALSO"):
        return False
    return None


def _lista(valor, separadores: str) -> list[str]:
    if valor is None:
        return []
    return [p.strip() for p in re.split(separadores, str(valor)) if p and p.strip()]


def _precio(valor) -> float | None:
    if valor is None or valor == "":
        return None
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        return None


def leer_maestro(data: bytes) -> Maestro:
    import openpyxl

    try:
        libro = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # archivo corrupto, no es xlsx...
        raise MaestroInvalido(f"No se pudo abrir el maestro: {e}") from e
    if "Maestro" not in libro.sheetnames:
        raise MaestroInvalido('El maestro debe tener una hoja llamada "Maestro".')

    hoja = libro["Maestro"]
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        raise MaestroInvalido("La hoja Maestro está vacía.")
    cab = [str(c).strip().lower() if c is not None else "" for c in filas[0]]
    faltan = [c for c in OBLIGATORIAS if c not in cab]
    if faltan:
        raise MaestroInvalido("Faltan columnas en el maestro: " + ", ".join(faltan))
    pos = {c: cab.index(c) for c in cab if c}

    registros = []
    for n, f in enumerate(filas[1:], start=2):
        def v(col):
            j = pos.get(col)
            return f[j] if j is not None and j < len(f) else None

        if all(x is None or str(x).strip() == "" for x in f):
            continue
        orden = v("orden")
        try:
            orden = float(orden)
        except (TypeError, ValueError):
            orden = 1e9
        registros.append(
            {
                "fila": n,
                "codigo": str(v("codigo") or "").strip(),
                "tipo": str(v("tipo") or "").strip(),
                "nombre": re.sub(r"\s+", " ", str(v("nombre") or "")).strip(),
                "descripcion": str(v("descripcion") or "").strip(),
                "categoria": str(v("categoria") or "").strip(),
                "modelos": _lista(v("modelos"), r"[,;]"),
                "codigos_sistema": _lista(v("codigo_sistema"), r"[;,\n]"),
                "precio_usd": _precio(v("precio_usd")),
                "controla_stock": _si_no(v("controla_stock")),
                "activo": _si_no(v("activo")),
                "imagenes_explicitas": [stem(x) for x in _lista(v("imagen"), r"[;,]")],
                "orden": orden,
                "observaciones": str(v("observaciones") or "").strip(),
            }
        )
    items = pd.DataFrame(registros)

    modelos: dict[str, str] = {}
    if "Modelos" in libro.sheetnames:
        for f in list(libro["Modelos"].iter_rows(values_only=True))[1:]:
            if f and f[0]:
                modelos[str(f[0]).strip()] = str(f[1]).strip() if len(f) > 1 and f[1] else ""
    return Maestro(items=items, modelos=modelos)


def validar_maestro(m: Maestro, imagenes_disponibles: set[str] | None = None) -> list[Hallazgo]:
    h: list[Hallazgo] = []
    it = m.items
    if it.empty:
        return [Hallazgo(ERROR, "Maestro", "Maestro", "El maestro no tiene filas.")]

    for cod, grupo in it.groupby("codigo"):
        if cod and len(grupo) > 1:
            h.append(Hallazgo(ERROR, "Maestro", cod, f"Código repetido en las filas {', '.join(map(str, grupo['fila']))}."))
    for _, r in it.iterrows():
        ref = r["codigo"] or f"fila {r['fila']}"
        if not r["codigo"]:
            h.append(Hallazgo(ERROR, "Maestro", ref, "Fila sin código."))
        if not r["nombre"]:
            h.append(Hallazgo(ERROR, "Maestro", ref, "Fila sin nombre."))
        if r["tipo"] not in TIPOS:
            h.append(Hallazgo(ERROR, "Maestro", ref, f"Tipo no válido: «{r['tipo']}». Use: {', '.join(TIPOS)}."))
        if r["controla_stock"] is None:
            h.append(Hallazgo(ERROR, "Maestro", ref, "controla_stock debe ser SI o NO."))
        if r["activo"] is None:
            h.append(Hallazgo(ERROR, "Maestro", ref, "activo debe ser SI o NO."))

    activos = it[it["activo"] == True]  # noqa: E712
    sin_precio = activos[activos["precio_usd"].isna() | (activos["precio_usd"] <= 0)]
    for _, r in sin_precio.iterrows():
        h.append(Hallazgo(ERROR, "Precio", r["codigo"], "Producto activo sin precio válido en USD."))
    marcador = activos[activos["precio_usd"] == PRECIO_MARCADOR]
    if len(marcador):
        h.append(Hallazgo(AVISO, "Precio", f"{len(marcador)} productos", f"Siguen con el precio de marcador de posición ({PRECIO_MARCADOR} USD)."))

    for _, r in activos[(activos["controla_stock"] == True) & (activos["codigos_sistema"].map(len) == 0)].iterrows():  # noqa: E712
        h.append(Hallazgo(AVISO, "Stock", r["codigo"], f"«{r['nombre']}» controla stock pero no tiene codigo_sistema: saldrá siempre sin stock."))
    for _, r in activos[(activos["tipo"] == "Accesorio") & (activos["modelos"].map(len) == 0)].iterrows():
        h.append(Hallazgo(AVISO, "Maestro", r["codigo"], f"«{r['nombre']}» no tiene modelos compatibles."))

    usados: dict[str, list[str]] = {}
    for _, r in activos.iterrows():
        for c in r["codigos_sistema"]:
            usados.setdefault(clave(c), []).append(r["codigo"])
    for k, cods in usados.items():
        if len(set(cods)) > 1:
            h.append(Hallazgo(AVISO, "Stock", ", ".join(sorted(set(cods))), "Comparten el mismo codigo_sistema: el stock se contará en ambos."))

    if imagenes_disponibles is not None:
        minusculas = {s.lower() for s in imagenes_disponibles}
        sin_foto = []
        for _, r in activos[activos["tipo"] == "Accesorio"].iterrows():
            tiene = any(s.lower() in minusculas for s in r["imagenes_explicitas"]) or any(
                re.fullmatch(re.escape(r["codigo"].lower()) + r"(_\d+)?", s) for s in minusculas
            )
            if not tiene:
                sin_foto.append(r["codigo"])
        if sin_foto:
            h.append(Hallazgo(AVISO, "Imágenes", f"{len(sin_foto)} accesorios", "Sin foto: " + ", ".join(sin_foto[:15]) + ("…" if len(sin_foto) > 15 else "")))
        faltan = []
        for _, r in activos.iterrows():
            for s in r["imagenes_explicitas"]:
                if s.lower() not in minusculas:
                    faltan.append(f"{r['codigo']} → {s}")
        if faltan:
            h.append(Hallazgo(AVISO, "Imágenes", f"{len(faltan)} referencias", "El maestro menciona archivos que no están en la carpeta: " + "; ".join(faltan[:8])))

    desconocidos = sorted({mo for ms in activos["modelos"] for mo in ms} - set(m.modelos))
    if m.modelos and desconocidos:
        h.append(Hallazgo(INFO, "Maestro", ", ".join(desconocidos), "Modelos usados que no figuran en la hoja Modelos (se mostrará solo el código)."))
    return h
